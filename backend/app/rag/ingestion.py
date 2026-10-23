from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional
from app.utils.text import clean_text, normalize_whitespace


@dataclass
class ParsedPage:
    text: str
    page_number: int
    section: Optional[str] = None


@dataclass
class ParsedDocument:
    pages: list[ParsedPage] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx", ".md"}
UPLOAD_DIR = Path("data/documents").resolve()


def parse_document(file_path: str, file_name: str) -> ParsedDocument:
    ext = Path(file_name).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")

    resolved = Path(file_path).resolve()
    if not str(resolved).startswith(str(UPLOAD_DIR)):
        raise ValueError("Invalid file path: outside upload directory")

    if not resolved.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if ext == ".pdf":
        return _parse_pdf(file_path)
    elif ext == ".docx":
        return _parse_docx(file_path)
    elif ext in (".txt", ".md"):
        return _parse_text(file_path)
    raise ValueError(f"Unsupported file type: {ext}")


def _parse_pdf(file_path: str) -> ParsedDocument:
    import pypdf

    doc = ParsedDocument(metadata={"source_type": "pdf"})
    reader = pypdf.PdfReader(file_path)
    for i, page in enumerate(reader.pages):
        raw = page.extract_text() or ""
        text = clean_text(normalize_whitespace(raw))
        if text.strip():
            doc.pages.append(ParsedPage(text=text, page_number=i + 1))
    if not doc.pages:
        raise ValueError("PDF appears to be empty or unreadable.")
    return doc


def _parse_docx(file_path: str) -> ParsedDocument:
    import docx

    document = docx.Document(file_path)
    doc = ParsedDocument(metadata={"source_type": "docx"})
    current_section: Optional[str] = None
    page_text_parts: list[str] = []

    for para in document.paragraphs:
        if para.style.name.startswith("Heading"):
            current_section = para.text.strip()
        text = para.text.strip()
        if text:
            page_text_parts.append(text)

    # DOCX has no page concept; treat whole doc as page 1
    full_text = clean_text("\n".join(page_text_parts))
    if not full_text:
        raise ValueError("DOCX appears to be empty.")
    doc.pages.append(ParsedPage(text=full_text, page_number=1, section=current_section))
    return doc


def _parse_text(file_path: str) -> ParsedDocument:
    resolved = Path(file_path).resolve()
    if not str(resolved).startswith(str(UPLOAD_DIR)):
        raise ValueError("Invalid file path: outside upload directory")
    with open(resolved, "r", encoding="utf-8", errors="replace") as f:
        raw = f.read()
    text = clean_text(raw)
    if not text:
        raise ValueError("Document appears to be empty.")
    doc = ParsedDocument(metadata={"source_type": "text"})
    doc.pages.append(ParsedPage(text=text, page_number=1))
    return doc
